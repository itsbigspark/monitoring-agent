#!/usr/bin/env python3
"""
Generate the client-facing Sentinel proposal slide deck (.pptx).
Content derived from incident-investigation-agent-proposal.md (client-safe:
excludes internal Appendix A go-to-market/commercial targeting).
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- palette ----
NAVY   = RGBColor(0x0F, 0x2A, 0x47)   # deep navy
TEAL   = RGBColor(0x18, 0x9A, 0xB4)   # accent teal
SLATE  = RGBColor(0x33, 0x3F, 0x4A)   # body text
LIGHT  = RGBColor(0xF2, 0xF5, 0xF7)   # light panel
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
MUTED  = RGBColor(0x6B, 0x77, 0x80)

prs = Presentation()
prs.slide_width  = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def rect(s, x, y, w, h, color, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=6, line_spacing=1.05):
    """runs: list of paragraphs; each paragraph is list of (text, size, bold, color, bullet_level)."""
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line_spacing
        if isinstance(para, tuple):
            para = [para]
        lvl = para[0][4] if len(para[0]) > 4 else 0
        p.level = lvl
        for seg in para:
            text, size, bold, color = seg[0], seg[1], seg[2], seg[3]
            r = p.add_run(); r.text = text
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = "Calibri"
    return tb

def header(s, kicker, title):
    rect(s, 0, 0, SW, Inches(1.35), NAVY)
    rect(s, 0, Inches(1.35), SW, Pt(4), TEAL)
    txt(s, Inches(0.6), Inches(0.18), Inches(12), Inches(0.4),
        [[(kicker.upper(), 12, True, TEAL)]])
    txt(s, Inches(0.6), Inches(0.5), Inches(12.1), Inches(0.8),
        [[(title, 28, True, WHITE)]])

def bullets(s, items, x=Inches(0.7), y=Inches(1.7), w=Inches(12), h=Inches(5.4),
            size=18, gap=10):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            text, lvl = it
        else:
            text, lvl = it, 0
        prefix = "•  " if lvl == 0 else "–  "
        color = SLATE if lvl == 0 else MUTED
        paras.append([(prefix + text, size - (lvl*2), False, color, lvl)])
    txt(s, x, y, w, h, paras, space_after=gap, line_spacing=1.05)

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

def chip(s, x, y, w, label, color=LIGHT, tcolor=NAVY, h=Inches(0.95), size=14, bold=True):
    rect(s, x, y, w, h, color)
    txt(s, x, y, w, h, [[(label, size, bold, tcolor)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 1 TITLE
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(5.35), SW, Pt(4), TEAL)
txt(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.4),
    [[("Sentinel", 54, True, WHITE)]])
txt(s, Inches(0.9), Inches(3.05), Inches(11.5), Inches(0.7),
    [[("Agentic incident auto-investigation", 26, False, TEAL)]])
txt(s, Inches(0.9), Inches(3.95), Inches(11.5), Inches(0.6),
    [[("The first feature of ", 18, False, RGBColor(0xC8,0xD4,0xDE)),
      ("Mission Control", 18, True, WHITE),
      (" — an application operations platform.", 18, False, RGBColor(0xC8,0xD4,0xDE))]])
txt(s, Inches(0.9), Inches(4.6), Inches(11.5), Inches(0.7),
    [[("Automating the investigation of production incidents — from alert to "
       "root cause and proposed fix.", 16, False, RGBColor(0xC8,0xD4,0xDE))]])
txt(s, Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.5),
    [[("Project proposal  ·  Draft for discussion", 13, False, MUTED)]])
notes(s, "One-line pitch: we turn a manual, senior-engineer-dependent investigation "
         "loop into a deployable agent that diagnoses incidents and proposes fixes. "
         "Sentinel is the first feature we ship of a broader platform, Mission Control. "
         "Working name 'Sentinel' — placeholder.")

# ---------------------------------------------------------------- 2 PROBLEM
s = slide()
header(s, "The problem", "An alert tells you something broke — not why, or how to fix it")
bullets(s, [
    "Splunk alerts fire off application logs and raise a ticket in ServiceNow.",
    "The ticket says \"this alert fired, here's the error message\" — and little else.",
    ("No root-cause analysis. No context. No proposed fix.", 1),
    "A senior engineer must then manually investigate every ticket.",
    "Investigation is the slowest, most repetitive phase of incident response — "
    "and it bottlenecks on a handful of people.",
], y=Inches(1.8))
notes(s, "Set up the pain: the alerting system detects symptoms but does no diagnosis. "
         "Humans carry the entire investigation load.")

# ---------------------------------------------------------------- 3 TODAY
s = slide()
header(s, "How it works today", "The current manual investigation loop")
steps = ["Alert fires\n(Splunk)", "Ticket raised\n(ServiceNow)",
         "Engineer reads\nlogs manually", "Cross-checks code\n& other systems",
         "Reproduces &\ndiagnoses", "Writes up fix\nfor dev team"]
n = len(steps); gap = Inches(0.18)
total = SW - Inches(1.4)
cw = Emu(int((total - gap*(n-1)) / n))
x = Inches(0.7); y = Inches(2.6)
for i, st in enumerate(steps):
    chip(s, x, y, cw, st, color=LIGHT, tcolor=NAVY, h=Inches(1.3), size=13)
    x = Emu(x + cw + gap)
txt(s, Inches(0.7), Inches(4.4), Inches(12), Inches(2),
    [[("Slow.", 18, True, NAVY), ("  Serial, manual evidence-gathering inflates "
      "mean-time-to-diagnosis.", 18, False, SLATE)],
     [("Person-dependent.", 18, True, NAVY), ("  Knowledge concentrated in a few "
      "senior engineers.", 18, False, SLATE)],
     [("Lossy.", 18, True, NAVY), ("  Findings live in tickets and people's heads; "
      "recurring patterns aren't captured.", 18, False, SLATE)]],
    space_after=10)
notes(s, "Walk the loop left to right. Emphasise that steps 3-5 are pure manual toil "
         "that the agent can take over.")

# ---------------------------------------------------------------- 4 INSIGHT
s = slide()
header(s, "Why now", "The loop is already proven — manually")
bullets(s, [
    "A custom Splunk MCP server (already built at NatWest) lets an AI assistant "
    "query logs on demand.",
    "Connected to the codebase, it already investigates incidents and surfaces "
    "root causes — it has been \"phenomenal\" in practice.",
    "Today this is run by hand, per incident. The opportunity is to automate and "
    "productise it.",
    "Timing: teams (e.g. Complaints) are onboarding the alerting system now — "
    "early enough to build automation in from the start.",
], y=Inches(1.9))
notes(s, "This de-risks the pitch: we're not proposing something speculative — we're "
         "automating a workflow that already works manually.")

# ---------------------------------------------------------------- 5 SOLUTION
s = slide()
header(s, "The proposed solution", "Sentinel — an agent that investigates incidents automatically")
bullets(s, [
    "Triggers automatically on a qualifying ServiceNow incident — no human kick-off.",
    "Pulls the relevant logs (Splunk), inspects the deployed codebase, and queries "
    "supporting systems (DB, Snowflake, S3) as needed.",
    "Reasons about the failure inside an isolated playground environment.",
    "Delivers a structured report: Overview → Investigation → Root Cause → "
    "Proposed Fix or Synopsis, with a confidence level.",
    "Posts findings back to the ticket / Teams; a human reviews and approves.",
], y=Inches(1.8))
notes(s, "This is the 'what'. Keep it outcome-focused; architecture comes next.")

# ---------------------------------------------------------------- 6 FLOW
s = slide()
header(s, "How Sentinel works", "End-to-end investigation flow")
flow = ["Detect &\ntriage", "Plan\ninvestigation", "Gather evidence\n(logs · code · data)",
        "Correlate &\nfind root cause", "Fix in playground\nor synopsise",
        "Human review\n& approve"]
n = len(flow); gap = Inches(0.16)
total = SW - Inches(1.4)
cw = Emu(int((total - gap*(n-1)) / n))
x = Inches(0.7); y = Inches(2.3)
for i, st in enumerate(flow):
    col = TEAL if i in (4,) else NAVY
    chip(s, x, y, cw, st, color=LIGHT if i!=4 else RGBColor(0xE4,0xF3,0xF6),
         tcolor=NAVY, h=Inches(1.4), size=13)
    x = Emu(x + cw + gap)
bullets(s, [
    "Iterative, cost-aware log retrieval — searches and narrows rather than dumping "
    "everything.",
    "Every diagnosis carries a confidence level and a full evidence trail.",
    "Every tool call and decision is logged for a complete audit trail.",
    "Human verdicts feed a golden dataset — tracking accuracy, fix acceptance and MTTD over time.",
], y=Inches(4.18), w=Inches(12), h=Inches(2.55), size=15.5, gap=7)
notes(s, "The playground step (highlighted) is where it validates a candidate fix "
         "safely before proposing it.")

# ---------------------------------------------------------------- 7 FIX OR SYNOPSIS
s = slide()
header(s, "Two valuable outcomes", "Fix it — or fast-track the human")
half = Inches(5.9)
# left card
rect(s, Inches(0.7), Inches(1.9), half, Inches(4.4), LIGHT)
rect(s, Inches(0.7), Inches(1.9), half, Inches(0.7), TEAL)
txt(s, Inches(0.9), Inches(1.9), half, Inches(0.7),
    [[("A — Validated fix", 18, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.95), Inches(2.8), Inches(5.4), Inches(3.4),
    [[("Reproduces the issue and develops a fix in an isolated playground.",16,False,SLATE)],
     [("Tests it before proposing — zero production risk.",16,False,SLATE)],
     [("Proposed for human review, e.g. as a pull request.",16,False,SLATE)]],
    space_after=12)
# right card
rx = Inches(6.75)
rect(s, rx, Inches(1.9), half, Inches(4.4), LIGHT)
rect(s, rx, Inches(1.9), half, Inches(0.7), NAVY)
txt(s, Emu(rx+Inches(0.2)), Inches(1.9), half, Inches(0.7),
    [[("B — Investigation synopsis", 18, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
txt(s, Emu(rx+Inches(0.25)), Inches(2.8), Inches(5.4), Inches(3.4),
    [[("When a reliable fix isn't possible, delivers what it found.",16,False,SLATE)],
     [("Overview, evidence, root-cause hypothesis, suggested next steps.",16,False,SLATE)],
     [("Gives the human a major head start — speeds up resolution.",16,False,SLATE)]],
    space_after=12)
txt(s, Inches(0.7), Inches(6.5), Inches(12), Inches(0.7),
    [[("Partial automation still wins: ~70% resolved + ~30% accelerated is a "
       "brilliant outcome.", 16, True, NAVY)]], align=PP_ALIGN.CENTER)
notes(s, "Key commercial framing: value does not require solving 100%. Even partial "
         "coverage plus acceleration of the rest is a strong ROI story.")

# ---------------------------------------------------------------- 8 ARCHITECTURE
s = slide()
header(s, "Architecture", "Built to scale and to swap parts out")
bullets(s, [
    "Agent orchestration (OpenAI Agents SDK) — a deterministic pipeline with agentic "
    "loops for evidence-gathering and fix-validation; full audit trace.",
    "MCP-first integration — reuses the existing Splunk MCP; each new system or "
    "incident type is a new connector + playbook, not a rewrite.",
    "Model-agnostic LLM — swap between Bedrock, self-hosted/in-VPC, or local "
    "(LM Studio / Ollama) by configuration. No vendor lock-in.",
    "Codebase enrichment (Intent Layer) — hierarchical, in-repo context so the "
    "agent understands the architecture before reading a line of code.",
    "Fully deployed inside the client environment — one isolated instance per client.",
    "Portable by design — the same versioned containers run on-prem or in any cloud via "
    "Compose / Helm; air-gap capable and built on open protocols.",
], y=Inches(1.72), size=16)
notes(s, "Three pillars to land: MCP-first (scales), model-agnostic (compliance/lock-in), "
         "in-client (data never leaves). Intent Layer is the quality multiplier. Engine note: "
         "Sentinel uses the OpenAI Agents SDK (right-sized; human gate is a post-run boundary, not "
         "a mid-run pause); LangGraph is kept for heavier future incident types, chosen per type "
         "behind shared seams.")

# ----------------------------------------------- 8b ARCHITECTURE DIAGRAM
# Companion visual. Requires sentinel-architecture.png (cropped) from
# build_architecture.py + crop step; skipped gracefully if absent.
import os as _os
_diagram = "/Users/zein/monitoring-agent/docs/diagrams/sentinel-architecture.png"
if _os.path.exists(_diagram):
    s = slide()
    header(s, "Architecture", "Proposed system architecture")
    from PIL import Image as _Image
    _cw, _ch = _Image.open(_diagram).size
    _atop, _abot = Inches(1.55), SH - Inches(0.25)
    _aw, _ah = SW - Inches(1.0), _abot - Inches(1.55)
    _asp = _cw / _ch
    _h = _ah; _w = Emu(int(_h * _asp))
    if _w > _aw:
        _w = _aw; _h = Emu(int(_w / _asp))
    _left = Emu(int((SW - _w) / 2)); _top = Emu(int(_atop + (_ah - _h) / 2))
    s.shapes.add_picture(_diagram, _left, _top, width=_w, height=_h)
    notes(s, "Companion visual to the architecture slide: ServiceNow -> ingestion/triage "
             "-> Agents SDK pipeline -> MCP layer -> read-only data sources; model-agnostic "
             "LLM + isolated playground on the right; all inside the client boundary; "
             "outputs go back to ticket / PR / Teams for human review.")

# ---------------------------------------------------------------- 9 SECURITY
s = slide()
header(s, "Built for the bank", "Security, compliance & governance — first-class")
cards = [
    ("Read-only, least privilege", "No write access to production. Scoped credentials."),
    ("Human-approved by default", "Fixes validated in a sandbox, proposed for review."),
    ("Fully in-client", "All logs, code, data & inference stay inside your boundary."),
    ("Complete audit trail", "Every tool call and decision logged and attributable."),
    ("Model choice you approve", "Run only models your security teams have signed off."),
    ("Sensitive-data handling", "Screening/redaction before anything reaches the model."),
]
cw2 = Inches(3.95); ch = Inches(1.9); gx = Inches(0.25); gy = Inches(0.3)
x0 = Inches(0.7); y0 = Inches(1.8)
for i, (t, d) in enumerate(cards):
    r, c = divmod(i, 3)
    x = Emu(x0 + c*(cw2+gx)); y = Emu(y0 + r*(ch+gy))
    rect(s, x, y, cw2, ch, LIGHT)
    rect(s, x, y, Pt(5), ch, TEAL)
    txt(s, Emu(x+Inches(0.2)), Emu(y+Inches(0.15)), Emu(cw2-Inches(0.3)), Inches(0.6),
        [[(t, 15, True, NAVY)]])
    txt(s, Emu(x+Inches(0.2)), Emu(y+Inches(0.75)), Emu(cw2-Inches(0.3)), Inches(1.0),
        [[(d, 12.5, False, SLATE)]])
notes(s, "This slide is what gets you past infosec. Lead with read-only + in-client + "
         "human-in-the-loop.")

# ------------------------------------------------------ 9b DEPLOYMENT & ONBOARDING
s = slide()
header(s, "Deployment & onboarding", "Designed to run anywhere — without client-specific forks")
txt(s, Inches(0.7), Inches(1.62), Inches(12), Inches(0.65),
    [[("One product, configured for each client and validated before go-live:",
       16, False, SLATE)]])
dcards = [
    ("Portable by default",
     "The same versioned container images run on-prem or in AWS, Azure or GCP — via "
     "Docker Compose or Helm, with an offline bundle for air-gapped environments."),
    ("Configuration, not custom code",
     "A deployment is a reviewed config bundle + secret references + MCP/content-source "
     "endpoints. Open protocols avoid cloud lock-in and per-client forks."),
    ("Preflight before go-live",
     "The onboarding CLI validates config, connectivity, documentation access and index "
     "freshness, then bootstraps and smoke-tests retrieval and an investigation."),
]
dn = 3; dgap = Inches(0.3); dtotal = SW - Inches(1.4)
dcw = Emu(int((dtotal - dgap*(dn-1)) / dn)); dx = Inches(0.7); dy = Inches(2.35)
for title, desc in dcards:
    rect(s, dx, dy, dcw, Inches(3.15), LIGHT)
    rect(s, dx, dy, dcw, Inches(0.78), NAVY)
    txt(s, dx, dy, dcw, Inches(0.78), [[(title, 15, True, WHITE)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Emu(dx+Inches(0.25)), Emu(dy+Inches(1.02)),
        Emu(dcw-Inches(0.5)), Inches(1.95), [[(desc, 13.2, False, SLATE)]])
    dx = Emu(dx + dcw + dgap)
rect(s, Inches(0.7), Inches(5.8), dtotal, Inches(0.72), RGBColor(0xE4,0xF3,0xF6))
txt(s, Inches(0.85), Inches(5.8), Emu(dtotal-Inches(0.3)), Inches(0.72),
    [[("One isolated instance per client · no phone-home · all code, evidence and inference stay inside the approved boundary",
       13.5, True, NAVY)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "This is the newly designed deployment model, not a claim that every packaging and "
         "onboarding component is already implemented. The key message is repeatability: shared "
         "versioned images, per-client configuration, automated dependency/documentation checks, "
         "and an air-gap path — not bespoke forks.")

# ---------------------------------------------------------------- 10 TIER 3
s = slide()
header(s, "Where we start", "Prove it on lower-tier systems first")
bullets(s, [
    "Begin on Tier 3 platforms — low criticality, ~3-day resolution turnaround.",
    "An instant automated investigation there turbocharges the monitoring teams' "
    "resolution-time KPIs and dashboards.",
    "Demonstrate value and build trust with zero risk to business-critical systems.",
    "Then scale up the tiers as confidence grows.",
], y=Inches(1.9))
notes(s, "Tier 1 needs instant human action; Tier 3 has slack we can fill — easy, "
         "low-risk, KPI-boosting first win.")

# ---------------------------------------------------------------- 11 SCALE
s = slide()
header(s, "How it scales", "One system today — the whole estate over time")
bullets(s, [
    "MCP-first design means new incident types and new systems are additive, "
    "reusing the core agent, governance and reporting unchanged.",
    "Start with Splunk-driven incidents on microservices; broaden to any incident "
    "type — e.g. Airflow DAG/task failures, or LLM-output quality-metric alerts.",
    "Triggers generalise too — ServiceNow today; Airflow events and email alerts next.",
    "Progressive autonomy — automation is earned: human-approved first, then "
    "opt-in auto-apply on lower-tier systems once the track record is proven "
    "(always audited and reversible).",
    "Applicable across independent business areas, each with many candidate systems.",
], y=Inches(1.85), gap=9)
notes(s, "Land the scalability story without internal targeting specifics. Progressive "
         "autonomy reassures: we don't ask for trust up front, we earn it. The Airflow DAG-failure "
         "and LLM-eval examples are especially apt for a data-science/AIOps buyer — and they also "
         "become monitoring surfaces in Mission Control (DAG stats, LLM eval-metric trends).")

# ----------------------------------------------- 11b MISSION CONTROL
s = slide()
header(s, "The bigger picture", "Sentinel is the first feature of Mission Control")
bullets(s, [
    "Mission Control is an operations cockpit for the teams that monitor and support "
    "many applications across a shared incident queue.",
    "It organises incidents per application, centralises monitoring dashboards, and — "
    "through Sentinel — auto-investigates and resolves them.",
], y=Inches(1.7), size=17, gap=8)
feats = [
    ("Sentinel", "Auto-investigation\n& resolution", True),
    ("Incident cockpit", "Per-application,\nhistorical view", False),
    ("Dashboard hub", "QuickSight, Tableau\nin one place", False),
    ("Ask Mission Control", "Docs + deployed code\nQ&A · future", False),
]
nf = 4; fgap = Inches(0.3); ftotal = SW - Inches(1.4)
fcw = Emu(int((ftotal - fgap*(nf-1)) / nf)); fx = Inches(0.7); fy = Inches(3.45)
for name, desc, first in feats:
    head_color = TEAL if first else NAVY
    rect(s, fx, fy, fcw, Inches(1.75), LIGHT)
    rect(s, fx, fy, fcw, Inches(0.6), head_color)
    txt(s, fx, fy, fcw, Inches(0.6), [[(name, 15, True, WHITE)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, fx, Emu(fy+Inches(0.72)), fcw, Inches(0.8),
        [[(desc, 12.5, False, SLATE)]], align=PP_ALIGN.CENTER)
    if first:
        txt(s, fx, Emu(fy+Inches(1.42)), fcw, Inches(0.3),
            [[("◆ ships first", 11, True, TEAL)]], align=PP_ALIGN.CENTER)
    fx = Emu(fx + fcw + fgap)
rect(s, Inches(0.7), Inches(5.45), ftotal, Inches(0.62), RGBColor(0xE4,0xF3,0xF6))
txt(s, Inches(0.7), Inches(5.45), ftotal, Inches(0.62),
    [[("Application catalog — shared foundation (apps · repos · log sources · required docs · playbooks)",
       13, True, NAVY)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.7), Inches(6.32), Inches(12), Inches(0.75),
    [[("ServiceNow owns incidents. Confluence owns documentation. Mission Control indexes, enriches and acts — it does not replace either.",
       13.2, True, MUTED)]], align=PP_ALIGN.CENTER)
notes(s, "Frame: we ship Sentinel first (the wedge and the differentiator). It is the first "
         "feature of a broader platform, Mission Control, for the monitoring teams that own "
         "many apps. Same in-client deployment. ServiceNow stays authoritative. The dashboard "
         "hub and analytics follow on the same application-catalog foundation. Keep the focus on "
         "Sentinel; this slide just shows where it's heading. Ask Mission Control is explicitly "
         "future scope: permission-aware Q&A over Confluence plus deployed code/Intent Layer and "
         "authorised incident history. ServiceNow and Confluence remain authoritative.")

# ----------------------------------------------- 11c OPERATIONAL KNOWLEDGE
s = slide()
header(s, "Operational knowledge", "Better handovers — without creating another source of truth")
txt(s, Inches(0.7), Inches(1.58), Inches(12), Inches(0.8),
    [[("Onboarding defines what every support team needs; Mission Control validates and indexes it for day-to-day use:",
       16, False, SLATE)]])
kcards = [
    ("Confluence stays authoritative",
     "Architecture, runbooks, ownership, dependencies, deployment/rollback and known failure "
     "modes are authored and governed in Confluence. Mission Control stores requirements, links, "
     "permissions and index freshness — not duplicate pages."),
    ("Support teams become self-sufficient",
     "Indexed documentation, code context and incident learning make routine questions and "
     "investigations answerable without repeatedly returning to the original developers. "
     "Engineering expertise stays focused on genuinely novel issues."),
    ("Future: code-aware Q&A",
     "Ask Mission Control will combine approved Confluence content with the relevant deployed "
     "code / Intent Layer and authorised incident history — a cited answer grounded in more than "
     "documentation alone."),
]
nk = 3; kgap = Inches(0.3); ktotal = SW - Inches(1.4)
kcw = Emu(int((ktotal - kgap*(nk-1)) / nk)); kx = Inches(0.7); ky = Inches(2.5)
for title, desc in kcards:
    rect(s, kx, ky, kcw, Inches(3.35), LIGHT)
    rect(s, kx, ky, kcw, Inches(0.85), NAVY)
    txt(s, Emu(kx+Inches(0.15)), ky, Emu(kcw-Inches(0.3)), Inches(0.85),
        [[(title, 14.5, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Emu(kx+Inches(0.23)), Emu(ky+Inches(1.02)), Emu(kcw-Inches(0.46)), Inches(2.15),
        [[(desc, 12.8, False, SLATE)]])
    kx = Emu(kx + kcw + kgap)
txt(s, Inches(0.7), Inches(6.12), Inches(12), Inches(0.85),
    [[("Outcome: a durable handover, fewer routine developer escalations, and application knowledge that improves with use.",
       14.5, True, NAVY)]], align=PP_ALIGN.CENTER)
notes(s, "Lead with structured handover and operational resilience, not criticism of current "
         "documentation. Confluence remains the source of truth; Mission Control defines the "
         "onboarding contract, checks accessibility/freshness and maintains a permission-aware "
         "index with citations. Say 'reduces routine dependence', not 'replaces developers'. "
         "Ask Mission Control is roadmap. The Rovo distinction is code context: planned answers "
         "combine Confluence with deployed code/Intent Layer and incident history, not docs alone.")

# ---------------------------------------------------------------- 12 VALUE
s = slide()
header(s, "The value", "Why this matters")
cards = [
    ("Faster resolution", "Compresses the slowest phase — time-to-diagnosis."),
    ("Senior-engineer leverage", "Frees scarce expertise from repetitive triage."),
    ("Measured quality", "Human verdicts + golden-set replay track accuracy and acceptance."),
    ("Support independence", "Structured handover reduces routine developer escalations."),
    ("Low-risk proving ground", "KPI gains on Tier 3 before critical systems."),
    ("Partial automation wins", "Resolve some, accelerate the rest — net big gain."),
]
cw2 = Inches(3.95); ch = Inches(1.9); gx = Inches(0.25); gy = Inches(0.3)
x0 = Inches(0.7); y0 = Inches(1.8)
for i, (t, d) in enumerate(cards):
    r, c = divmod(i, 3)
    x = Emu(x0 + c*(cw2+gx)); y = Emu(y0 + r*(ch+gy))
    rect(s, x, y, cw2, ch, LIGHT)
    rect(s, x, y, cw2, Pt(5), TEAL)
    txt(s, Emu(x+Inches(0.2)), Emu(y+Inches(0.2)), Emu(cw2-Inches(0.3)), Inches(0.6),
        [[(t, 15, True, NAVY)]])
    txt(s, Emu(x+Inches(0.2)), Emu(y+Inches(0.8)), Emu(cw2-Inches(0.3)), Inches(1.0),
        [[(d, 12.5, False, SLATE)]])
notes(s, "Baseline MTTD, accuracy, fix acceptance, senior-engineer hours and routine developer "
         "escalations before rollout so the after delta is provable. Human verdict capture and "
         "golden-dataset replay answer the buyer question: 'how do we know it's right?'")

# ---------------------------------------------------------------- 13 ENGAGEMENT
s = slide()
header(s, "Engagement model", "Low-commitment start, scalable rollout")
n = 3; gap = Inches(0.3)
total = SW - Inches(1.4)
cw = Emu(int((total - gap*(n-1)) / n))
stages = [
    ("1 · Platform analysis", "A short, fixed-fee analysis of a target platform — "
     "scope evidence access, documentation readiness, security and success baselines up front."),
    ("2 · Proof of Concept", "Configure Sentinel, preflight every dependency and demonstrate it "
     "on real or historical incidents — proof before commitment."),
    ("3 · Scaled rollout", "Implement platform by platform (fixed-price or T&M). "
     "Shared components built once; only platform-specific work per system."),
]
x = Inches(0.7); y = Inches(2.1)
for t, d in stages:
    rect(s, x, y, cw, Inches(3.4), LIGHT)
    rect(s, x, y, cw, Inches(0.8), NAVY)
    txt(s, x, y, cw, Inches(0.8), [[(t, 16, True, WHITE)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Emu(x+Inches(0.25)), Emu(y+Inches(1.0)), Emu(cw-Inches(0.5)), Inches(2.3),
        [[(d, 14, False, SLATE)]])
    x = Emu(x + cw + gap)
txt(s, Inches(0.7), Inches(5.9), Inches(12), Inches(1),
    [[("The product stays shared: each rollout adds a reviewed config bundle, approved knowledge sources, connectors and playbooks — not a fork.",
       13.5, False, MUTED)]], align=PP_ALIGN.CENTER)
notes(s, "Client-safe commercial framing. Keep specific numbers for the live "
         "conversation. The analysis-first step de-risks pricing for both sides.")

# ---------------------------------------------------------------- 14 NEXT STEPS
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(1.35), SW, Pt(4), TEAL)
txt(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.8),
    [[("Next steps", 30, True, WHITE)]])
txt(s, Inches(0.7), Inches(2.0), Inches(12), Inches(4),
    [[("1.  Align on a target platform for a Proof of Concept.", 20, False, WHITE)],
     [("2.  Run the fixed-fee platform analysis to scope & price it.", 20, False, WHITE)],
     [("3.  Deliver the PoC on real incidents and review the results.", 20, False, WHITE)],
     [("4.  Agree the rollout plan across systems and areas.", 20, False, WHITE)]],
    space_after=16, line_spacing=1.1)
txt(s, Inches(0.7), Inches(6.2), Inches(12), Inches(0.5),
    [[("Let's pick one platform and prove it.", 18, True, TEAL)]])
txt(s, Inches(0.7), Inches(6.72), Inches(12), Inches(0.4),
    [[("Sentinel first — the foundation for Mission Control.", 13, False, RGBColor(0xC8,0xD4,0xDE))]])
notes(s, "Close with a concrete, low-commitment ask: agree one platform and run the "
         "analysis. Momentum over perfection. Sentinel is step one toward Mission Control.")

out = "/Users/zein/monitoring-agent/docs/slide-deck/sentinel-proposal-deck.pptx"
prs.save(out)
print("saved", out, "slides:", len(prs.slides._sldIdLst))
