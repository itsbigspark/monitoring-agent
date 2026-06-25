#!/usr/bin/env python3
"""Crop the architecture PNG to content, then insert it as a new slide after
slide 8 of the existing deck (in place, preserving other slides)."""
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import copy

NAVY=RGBColor(0x0F,0x2A,0x47); TEAL=RGBColor(0x18,0x9A,0xB4)
WHITE=RGBColor(0xFF,0xFF,0xFF); MUTED=RGBColor(0x6B,0x77,0x80)

DOCS="/Users/zein/monitoring-agent/docs"
SRC=f"{DOCS}/sentinel-architecture.svg.png"
CROP=f"{DOCS}/sentinel-architecture.png"
DECK=f"{DOCS}/sentinel-proposal-deck.pptx"

# ---- crop white margins ----
im=Image.open(SRC).convert("RGB")
bg=Image.new("RGB",im.size,(255,255,255))
diff=ImageChops.difference(im,bg)
bbox=diff.getbbox()
pad=18
l,t,r,b=bbox
l=max(0,l-pad); t=max(0,t-pad); r=min(im.width,r+pad); b=min(im.height,b+pad)
im.crop((l,t,r,b)).save(CROP)
cw,ch=(r-l),(b-t)
print("cropped to",cw,"x",ch)

# ---- open deck, add slide ----
prs=Presentation(DECK)
SW,SH=prs.slide_width,prs.slide_height
s=prs.slides.add_slide(prs.slide_layouts[6])

def rect(x,y,w,h,color):
    sp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h)
    sp.fill.solid(); sp.fill.fore_color.rgb=color; sp.line.fill.background()
    sp.shadow.inherit=False; return sp
def txt(x,y,w,h,text,size,bold,color,align=PP_ALIGN.LEFT):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.size=Pt(size); r.font.bold=bold
    r.font.color.rgb=color; r.font.name="Calibri"

# header band (matches other slides)
rect(0,0,SW,Inches(1.35),NAVY)
rect(0,Inches(1.35),SW,Pt(4),TEAL)
txt(Inches(0.6),Inches(0.18),Inches(12),Inches(0.4),"ARCHITECTURE",12,True,TEAL)
txt(Inches(0.6),Inches(0.5),Inches(12.1),Inches(0.8),
    "Proposed system architecture",28,True,WHITE)

# place image, fit into area below header
area_top=Inches(1.55); area_bottom=SH-Inches(0.25)
avail_w=SW-Inches(1.0); avail_h=area_bottom-area_top
aspect=cw/ch
# fit by height first
h=avail_h; w=Emu(int(h*aspect))
if w>avail_w:
    w=avail_w; h=Emu(int(w/aspect))
left=Emu(int((SW-w)/2)); top=Emu(int(area_top+(avail_h-h)/2))
s.shapes.add_picture(CROP,left,top,width=w,height=h)
s.notes_slide.notes_text_frame.text=(
    "Companion visual to the architecture slide. Walk: ServiceNow -> ingestion/triage "
    "-> LangGraph investigation graph -> MCP layer -> read-only data sources; "
    "model-agnostic LLM + isolated playground on the right; everything inside the "
    "client environment boundary; outputs go back to the ticket / PR / Teams for "
    "human review.")

# ---- move new slide to position right after slide 8 (0-based index 8) ----
lst=prs.slides._sldIdLst
ids=list(lst)
new=ids[-1]
lst.remove(new)
lst.insert(8, new)   # insert as the 9th slide
prs.save(DECK)
print("deck now has", len(prs.slides._sldIdLst), "slides; diagram is slide 9")
